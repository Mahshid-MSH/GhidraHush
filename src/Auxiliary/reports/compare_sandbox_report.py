"""Excel reporting for GhidraHush binary-comparison results.

Owns Results/comparisons.xlsx: creates it (with a header row and basic
formatting) the first time it's touched, and appends one row per compared
binary pair on every subsequent call.

Each call to add_result() reopens the workbook from disk, appends, and
saves -- it doesn't hold the file open across calls. That makes results
survive a crash mid-run (you only ever lose the row in progress, not the
whole file) and makes it safe to call from separate `compare_sandbox.py`
invocations, not just within one long batch run. It is NOT safe against
two processes writing at the exact same instant -- last save wins, the
earlier one's row can be lost. If you expect concurrent runs against the
same workbook, serialize them (a lock file, or just don't run --manifest
jobs in parallel against the same Results folder).

No formulas are written here on purpose: openpyxl writes formulas as
strings with no cached value, so anything that reads this file before a
LibreOffice recalc pass sees None for those cells. That's an awkward
dependency for a log file a script appends to on every run, so every
column here is a plain literal value instead.
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SHEET_NAME = "LCS Results"

# (column header, key read from the dict passed to add_result())
COLUMNS = [
    ("Timestamp",             "timestamp"),
    ("Original Path",         "original_path"),
    ("Candidate Path",        "candidate_path"),
    ("CAPE Task (Original)",  "task_original"),
    ("CAPE Task (Candidate)", "task_candidate"),
    ("API Calls (Original)",  "api_calls_original"),
    ("API Calls (Candidate)", "api_calls_candidate"),
    ("LCS (raw)",             "lcs_raw"),
    ("LCS (collapsed)",       "lcs_collapsed"),
]

BODY_FONT = Font(name="Arial", size=10)
HEADER_FONT = Font(name="Arial", size=10, bold=True)
SCORE_COLUMNS = {"LCS (raw)", "LCS (collapsed)"}  # get a 0.0000 number format


class ComparisonReport:
    """Wraps Results/comparisons.xlsx. Cheap to construct repeatedly --
    all state lives in the file on disk, not on the instance."""

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            self._create_workbook()

    def _create_workbook(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        self._write_header(ws)
        wb.save(self.path)

    @staticmethod
    def _write_header(ws) -> None:
        for col_idx, (header, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(header) + 2)
        ws.freeze_panes = "A2"

    def add_result(self, row: dict) -> int:
        """Append one row of comparison results.

        `row` should carry the keys listed in COLUMNS (second element of
        each pair) -- a missing key is written blank, not an error, so you
        can pass the exact dict process_pair() already builds. `timestamp`
        is filled in automatically if not supplied.

        Returns the 1-indexed sheet row number written to.
        """
        wb = load_workbook(self.path)
        if SHEET_NAME not in wb.sheetnames:
            # sheet got renamed/removed by hand -- recreate rather than
            # silently writing to the wrong place
            ws = wb.create_sheet(SHEET_NAME)
            self._write_header(ws)
        else:
            ws = wb[SHEET_NAME]

        row = dict(row)
        row.setdefault("timestamp", _dt.datetime.now().isoformat(timespec="seconds"))

        ws.append([row.get(key, "") for _, key in COLUMNS])
        written_row = ws.max_row

        for col_idx, (header, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=written_row, column=col_idx)
            cell.font = BODY_FONT
            if header in SCORE_COLUMNS and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000"

        wb.save(self.path)
        return written_row