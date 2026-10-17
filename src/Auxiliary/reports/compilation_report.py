# Add this method inside your ComparisonReport class (or equivalent) in reporting/excel_report.py:

def add_compilation_stats(self, stats_list):
    """
    Appends or writes compilation summary stats to a separate sheet named 'Compilation Stats'.
    """
    import openpyxl
    from pathlib import Path

    wb = openpyxl.load_workbook(self.path) if Path(self.path).exists() else openpyxl.Workbook()
    sheet_name = "Compilation Stats"
    
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(title=sheet_name)
        # Write headers
        ws.append(["Binary Name", "Total Functions", "Compiled Successfully", "Failed Functions", "Total Repairs/Patches", "Failure %"])

    for item in stats_list:
        success = item.get("success", 0)
        failed = item.get("failed", 0)
        total = success + failed
        fail_pct = (failed / total * 100) if total > 0 else 0.0
        
        ws.append([
            item.get("binary", "Unknown"),
            total,
            success,
            failed,
            item.get("total_repairs", 0),
            round(fail_pct, 2)
        ])
    
    wb.save(self.path)