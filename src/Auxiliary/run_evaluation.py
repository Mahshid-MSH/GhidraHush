#!/usr/bin/env python3
"""Evaluate candidate executables with CAPE and VirusTotal."""
from __future__ import annotations

import argparse
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

if __package__:
    from . import compare_sandbox, VT_scanner
else:  # Direct script execution puts this directory on sys.path.
    import compare_sandbox
    import VT_scanner

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = PROJECT_ROOT / "Results" / "comparisons.xlsx"
EXECUTABLE_EXTENSIONS = {".exe", ".dll", ".sys"}
RUN_HEADERS = (
    "Run ID", "Artifact Type", "Executable", "CAPE Comparison",
    "VT Detections", "Status", "Details",
)


def prepare_executable_dirs(
    workspace_dir: str | Path,
    original_exe: str | Path | None = None,
) -> dict[str, Path]:
    """Create evaluation folders and optionally copy in the original binary."""
    root = Path(workspace_dir).resolve() / "executable_results"
    paths = {
        "original": root / "original",
        "diversified_binaries": root / "diversified_binaries",
    }
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    if original_exe is not None:
        source = Path(original_exe).resolve()
        if source.is_file():
            destination = paths["original"] / source.name
            if source != destination.resolve():
                shutil.copy2(source, destination)
    return paths


def _find_executables(directory: Path) -> list[Path]:
    return sorted(
        path for path in directory.rglob("*")
        if path.is_file() and path.suffix.lower() in EXECUTABLE_EXTENSIONS
    )


def _append_run_sheet(
    xlsx_path: Path,
    run_id: str,
    rows: list[dict[str, Any]],
) -> str:
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    if xlsx_path.exists():
        workbook = load_workbook(xlsx_path)
    else:
        workbook = Workbook()
        if workbook.active and workbook.active.title == "Sheet":
            workbook.remove(workbook.active)

    base_name = f"Evaluation {run_id}"[:31]
    sheet_name = base_name
    suffix = 1
    while sheet_name in workbook.sheetnames:
        tail = f"_{suffix}"
        sheet_name = f"{base_name[:31 - len(tail)]}{tail}"
        suffix += 1

    sheet = workbook.create_sheet(sheet_name)
    sheet.append(RUN_HEADERS)
    for row in rows:
        sheet.append([
            run_id,
            row.get("type", ""),
            row.get("path", ""),
            row.get("comparison", ""),
            row.get("vt_detections", ""),
            row.get("status", ""),
            row.get("details", ""),
        ])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 64
    sheet.column_dimensions["D"].width = 22
    sheet.column_dimensions["E"].width = 16
    sheet.column_dimensions["F"].width = 20
    sheet.column_dimensions["G"].width = 64
    workbook.save(xlsx_path)
    return sheet_name


def evaluate_workspace(
    workspace_dir: str | Path,
    original_exe: str | Path,
    xlsx_path: str | Path = DEFAULT_XLSX,
) -> list[dict[str, Any]]:
    """Run CAPE comparisons and VirusTotal scans for available candidates."""
    load_dotenv()
    folders = prepare_executable_dirs(workspace_dir, original_exe)
    source_original = Path(original_exe).resolve()
    original = folders["original"] / source_original.name
    candidates = _find_executables(folders["diversified_binaries"])
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    run_id = f"{Path(workspace_dir).resolve().name}_{timestamp}"
    outcomes: list[dict[str, Any]] = []

    if not candidates:
        outcomes.append({
            "status": "No candidates",
            "details": "Add candidate executables under executable_results/diversified_binaries.",
        })
        results_dir = Path(workspace_dir).resolve() / "executable_results"
        print(
            "[evaluation] no diversified executables found; "
            f"folders are ready under {results_dir}"
        )
    elif not original.is_file():
        outcomes.append({"status": "Original missing", "details": str(source_original)})
    else:
        api_key = os.environ.get("VT_API_KEY")
        for candidate in candidates:
            row: dict[str, Any] = {
                "type": "diversified",
                "path": str(candidate),
                "status": "Completed",
                "details": "",
            }
            try:
                compare_sandbox.process_pair(str(original), str(candidate), str(xlsx_path))
                row["comparison"] = "CAPE comparison saved"
            except Exception as exc:
                row["comparison"] = "Failed"
                row["details"] = f"CAPE: {type(exc).__name__}: {exc}"
                row["status"] = "Partial"
                print(f"[evaluation] CAPE comparison failed for {candidate}: {exc}")

            if api_key:
                try:
                    row["vt_detections"] = VT_scanner.scan_file(str(candidate), api_key)
                except Exception as exc:
                    row["status"] = "Partial"
                    extra = f"VirusTotal: {type(exc).__name__}: {exc}"
                    row["details"] = "; ".join(filter(None, (row["details"], extra)))
                    print(f"[evaluation] VirusTotal scan failed for {candidate}: {exc}")
            else:
                row["details"] = "; ".join(filter(
                    None,
                    (row["details"], "VirusTotal skipped: VT_API_KEY is not set"),
                ))
                if row["status"] == "Completed":
                    row["status"] = "CAPE only"
            outcomes.append(row)

        if api_key:
            try:
                VT_scanner.export_excel(str(xlsx_path), manifest_path="")
            except Exception as exc:
                outcomes.append({
                    "status": "Partial",
                    "details": f"VirusTotal workbook export failed: {type(exc).__name__}: {exc}",
                })
                print(f"[evaluation] VirusTotal workbook export failed: {exc}")

    sheet_name = _append_run_sheet(Path(xlsx_path), run_id, outcomes)
    print(f"[evaluation] run results saved to '{sheet_name}' in {xlsx_path}")
    return outcomes


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", required=True, help="Workspace run directory")
    parser.add_argument("--original", required=True, help="Original executable for comparison")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="Shared results workbook")
    args = parser.parse_args()
    evaluate_workspace(args.workspace, args.original, args.xlsx)


if __name__ == "__main__":
    main()
