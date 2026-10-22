#!/usr/bin/env python3
"""Compare CAPE reports and API-call sequences for two samples.

The normalized LCS score is LCS(original, candidate) / len(original).
Results are appended to Results/comparisons.xlsx.

Usage:
    python compare_sandbox.py original.exe recompiled.exe   # single pair
    python compare_sandbox.py --manifest pairs.csv           # batch
"""
import argparse
import csv
import os
import re
import sys
import time
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

import requests
import openpyxl

NOISE_APIS = {"NtTestAlert", "LdrpCallInitRoutine", "RtlUserThreadStart"}
HEX = re.compile(r"0x[0-9a-fA-F]+")

CAPE = os.environ.get("CAPE_API_URL", "http://host.docker.internal:8000/apiv2")
TOKEN = os.environ.get("CAPE_API_TOKEN", "")
HEADERS = {"Authorization": f"Token {TOKEN}"} if TOKEN else {}
OPTIONS = {"timeout": 60, "machine": "", "package": ""}
MAX_LCS_CELLS = 25_000_000
COMPARISON_HEADERS = (
    "Original Path", "Candidate Path", "Task Original", "Task Candidate",
    "API Calls Original", "API Calls Candidate", "LCS Raw", "LCS Collapsed",
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX_PATH = PROJECT_ROOT / "Results" / "comparisons.xlsx"


def append_to_excel(row, xlsx_path=DEFAULT_XLSX_PATH):
    """Append one comparison result to the workbook."""
    path = Path(xlsx_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()

    sheet_name = "Comparison Results"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active
        ws.title = sheet_name
        ws.append(COMPARISON_HEADERS)

    ws.append([
        row.get("original_path"),
        row.get("candidate_path"),
        row.get("task_original"),
        row.get("task_candidate"),
        row.get("api_calls_original"),
        row.get("api_calls_candidate"),
        row.get("lcs_raw"),
        row.get("lcs_collapsed"),
    ])

    wb.save(path)
    return path


def submit(path):
    with open(path, "rb") as f:
        response = requests.post(
            f"{CAPE}/tasks/create/file/",
            files={"file": (Path(path).name, f)},
            data=OPTIONS,
            headers=HEADERS,
        )
    response.raise_for_status()
    result = response.json()
    if result.get("error"):
        raise RuntimeError(result)
    return result["data"]["task_ids"][0]


def wait(task_id, poll=10, limit=1800):
    start = time.time()
    while time.time() - start < limit:
        response = requests.get(f"{CAPE}/tasks/status/{task_id}/", headers=HEADERS)
        status = response.json().get("data")
        print(f"[task {task_id}] {status}", file=sys.stderr)
        if status == "reported":
            return
        if status and status.startswith("failed"):
            raise RuntimeError(f"task {task_id} {status}")
        time.sleep(poll)
    raise TimeoutError(task_id)


def report(task_id):
    response = requests.get(
        f"{CAPE}/tasks/get/report/{task_id}/json/",
        headers=HEADERS,
    )
    response.raise_for_status()
    return response.json()


def features(rep):
    summ = rep.get("behavior", {}).get("summary", {})
    net = rep.get("network", {})
    feature_sets = {
        "signatures": {s["name"] for s in rep.get("signatures", [])},
        "mutexes": set(summ.get("mutexes", [])),
        "files_touched": set(summ.get("files", [])),
        "files_written": set(summ.get("write_files", [])),
        "regkeys_written": set(summ.get("write_keys", [])),
        "commands": set(summ.get("executed_commands", [])),
        "resolved_apis": set(summ.get("resolved_apis", [])),
        "domains": {d.get("domain") for d in net.get("domains", [])},
        "hosts": {h.get("ip") for h in net.get("hosts", [])},
        "http_uris": {h.get("uri") for h in net.get("http", [])},
        "dropped_sha256": {d.get("sha256") for d in rep.get("dropped", [])},
        "cape_payloads": {
            payload.get("sha256")
            for payload in rep.get("CAPE", {}).get("payloads", [])
            if isinstance(payload, dict)
        },
    }
    pe = (
        rep.get("target", {}).get("file", {}).get("pe")
        or rep.get("static", {}).get("pe", {})
    )
    imps = set()
    imports = pe.get("imports") or []
    dlls = imports.values() if isinstance(imports, dict) else imports
    for dll in dlls:
        for imported_function in dll.get("imports", []):
            imps.add(f'{dll.get("dll")}!{imported_function.get("name")}')
    feature_sets["pe_imports"] = imps
    return feature_sets


def compare(features_a, features_b):
    out = {}
    for key in features_a:
        out[key] = {
            "common": sorted(value for value in features_a[key] & features_b[key] if value),
            "only_A": sorted(value for value in features_a[key] - features_b[key] if value),
            "only_B": sorted(value for value in features_b[key] - features_a[key] if value),
        }
    return out


def sample_names(rep):
    n = rep.get("target", {}).get("file", {}).get("name", "")
    return sorted({n, n.rsplit(".", 1)[0]} - {""}, key=len, reverse=True)


def normalize(value, names):
    s = str(value)
    for n in names:
        s = re.sub(re.escape(n), "<SAMPLE>", s, flags=re.I)
    s = re.sub(r"(?i)\\Users\\[^\\]+", r"\\Users\\<USER>", s)
    return HEX.sub("<HEX>", s)


def call_features(rep):
    names = sample_names(rep)
    apis = Counter()
    categories = Counter()
    arguments = set()
    sequence = []
    raw_sequence = []
    processes = set()
    for p in rep.get("behavior", {}).get("processes", []):
        processes.add(normalize(p.get("process_name", ""), names))
        for c in p.get("calls", []):
            api = c.get("api")
            if not api or api in NOISE_APIS:
                continue
            apis[api] += 1
            categories[c.get("category", "?")] += 1
            raw_sequence.append(api)
            if not sequence or sequence[-1] != api:
                sequence.append(api)
            for a in c.get("arguments", []):
                v = a.get("value")
                if isinstance(v, str) and len(v) > 3 and not HEX.fullmatch(v):
                    arguments.add(f'{api}.{a.get("name")}={normalize(v, names)}')
    return {
        "apis": apis,
        "categories": categories,
        "args": arguments,
        "seq": sequence,
        "raw_seq": raw_sequence,
        "procs": processes,
    }


def print_call_diff(ca, cb, top=30):
    count_a = sum(ca["apis"].values())
    count_b = sum(cb["apis"].values())
    print(f"\n== API calls: {count_a} in A, {count_b} in B")
    api_names = set(ca["apis"]) | set(cb["apis"])
    rows = sorted(
        ((name, ca["apis"][name], cb["apis"][name]) for name in api_names),
        key=lambda row: abs(row[1] - row[2]),
        reverse=True,
    )
    print(f"{'API':38}{'A':>7}{'B':>7}{'delta':>8}")
    for name, count_a, count_b in [row for row in rows if row[1] != row[2]][:top]:
        print(f"{name:38}{count_a:>7}{count_b:>7}{count_b - count_a:>+8}")
    print("\nAPIs only in A:", sorted(set(ca["apis"]) - set(cb["apis"])))
    print("APIs only in B:", sorted(set(cb["apis"]) - set(ca["apis"])))
    print("Category counts A:", dict(ca["categories"]))
    print("Category counts B:", dict(cb["categories"]))
    sim = SequenceMatcher(
        None, ca["seq"][:5000], cb["seq"][:5000], autojunk=False
    ).ratio()
    print(f"\nAPI sequence similarity (difflib, collapsed, capped@5000): {sim:.1%}")
    for label, arguments, other_arguments in (
        ("A", ca["args"], cb["args"]),
        ("B", cb["args"], ca["args"]),
    ):
        only = sorted(arguments - other_arguments)
        print(f"\n== call arguments only in {label} ({len(only)}), first 40:")
        for line in only[:40]:
            print("  ", line)


def lcs_length(sequence_a, sequence_b):
    if len(sequence_a) > len(sequence_b):
        sequence_a, sequence_b = sequence_b, sequence_a
    n = len(sequence_a)
    if n == 0:
        return 0
    masks = {}
    for index, value in enumerate(sequence_a):
        masks[value] = masks.get(value, 0) | (1 << index)
    full = (1 << n) - 1
    v = full
    for value in sequence_b:
        mask = masks.get(value, 0)
        common = v & mask
        v = ((v + common) | (v & ~mask)) & full
    return n - bin(v).count("1")


def bounded_lcs_length(sequence_a, sequence_b, max_cells=MAX_LCS_CELLS):
    """Compatibility wrapper; the exact algorithm no longer needs a cell limit."""
    del max_cells
    return lcs_length(sequence_a, sequence_b)


def lcs_hat(seq_ref, seq_cand):
    if not seq_ref:
        return 0.0
    return bounded_lcs_length(seq_ref, seq_cand) / len(seq_ref)


def print_lcs(c_ref, c_cand):
    print("\n== Functionality Preservation (LCS), reference=original, candidate=recompiled")
    scores = {}
    fields = (("raw_seq", "raw", "lcs_raw"), ("seq", "collapsed", "lcs_collapsed"))
    for field, tag, key in fields:
        score = lcs_hat(c_ref[field], c_cand[field])
        scores[key] = score
        print(f"  lcs_hat [{tag:9}] = {score:.4f}  "
              f"(ref len={len(c_ref[field])}, cand len={len(c_cand[field])})")
    return scores


def process_pair(path_ref, path_cand, results_xlsx=DEFAULT_XLSX_PATH):
    """Submit, wait, fetch, diff and LCS-score one pair, then append to Excel."""
    t_ref, t_cand = submit(path_ref), submit(path_cand)
    print(
        f"submitted: original={t_ref} ({path_ref})  "
        f"candidate={t_cand} ({path_cand})",
        file=sys.stderr,
    )
    wait(t_ref)
    wait(t_cand)
    r_ref, r_cand = report(t_ref), report(t_cand)

    result = compare(features(r_ref), features(r_cand))
    for k, v in result.items():
        print(
            f"\n== {k}: {len(v['common'])} common, "
            f"{len(v['only_A'])} only in original, "
            f"{len(v['only_B'])} only in candidate"
        )

    c_ref, c_cand = call_features(r_ref), call_features(r_cand)
    print_call_diff(c_ref, c_cand)
    lcs_scores = print_lcs(c_ref, c_cand)

    row = {
        "original_path": path_ref, "candidate_path": path_cand,
        "task_original": t_ref, "task_candidate": t_cand,
        "api_calls_original": sum(c_ref["apis"].values()),
        "api_calls_candidate": sum(c_cand["apis"].values()),
        **lcs_scores,
    }

    xlsx_path = append_to_excel(row, results_xlsx)
    print(f"[xlsx] appended result to {xlsx_path}", file=sys.stderr)

    return row


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("sample_a", nargs="?", help="original binary")
    ap.add_argument("sample_b", nargs="?", help="candidate binary")
    ap.add_argument("--manifest", help="CSV, one 'original_path,candidate_path' row per pair")
    ap.add_argument("--summary-out", default="lcs_summary.csv",
                    help="CSV path for per-pair results in manifest mode")
    ap.add_argument("--results-xlsx", default=str(DEFAULT_XLSX_PATH),
                     help=f"path to the running Excel log (default: {DEFAULT_XLSX_PATH})")
    args = ap.parse_args()

    if args.manifest:
        rows = []
        with open(args.manifest, newline="") as f:
            for row in csv.reader(f):
                if not row or row[0].startswith("#"):
                    continue
                path_ref, path_cand = row[0].strip(), row[1].strip()
                print(f"\n{'='*70}\n{path_ref}  vs  {path_cand}\n{'='*70}")
                try:
                    rows.append(process_pair(path_ref, path_cand, args.results_xlsx))
                except Exception as e:
                    print(f"[!] pair failed: {path_ref} / {path_cand}: {e}", file=sys.stderr)
                    rows.append({
                        "original_path": path_ref,
                        "candidate_path": path_cand,
                        "error": str(e),
                    })
        if rows:
            with open(args.summary_out, "w", newline="") as f:
                fieldnames = sorted({key for row in rows for key in row})
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)
            print(f"\nwrote {len(rows)} pair(s) to {args.summary_out}")
    else:
        if not (args.sample_a and args.sample_b):
            ap.error("give two paths (original, candidate), or --manifest pairs.csv")
        process_pair(args.sample_a, args.sample_b, args.results_xlsx)


if __name__ == "__main__":
    main()