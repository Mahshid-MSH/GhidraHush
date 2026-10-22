import os
from datetime import datetime, timezone
from typing import Any
from openpyxl import Workbook, load_workbook

HEADERS = (
    "Run ID", "Binary", "Function", "Status", "Repairs", "Language",
    "Total Binary Success", "Total Binary Failed", "Total Repairs",
)


class CompilationReport:
    def __init__(self, file_path: str, sheet_name: str = "Compilation Stats") -> None:
        self.file_path = file_path
        self.sheet_name = sheet_name

    def add_compilation_stats(self, stats_list: list[dict[str, Any]]) -> None:
        """Append per-function results and one summary row per binary."""
        rows: list[dict[str, Any]] = []
        for entry in stats_list:
            binary = entry.get("binary", "")
            success = entry.get("success", 0)
            failed = entry.get("failed", 0)
            repairs = entry.get("total_repairs", 0)
            run_id = datetime.now(timezone.utc).isoformat(timespec="seconds")
            for function in entry.get("functions", []):
                rows.append({
                    "Run ID": run_id,
                    "Binary": binary,
                    "Function": function.get("name"),
                    "Status": "Success" if function.get("ok") else "Failed",
                    "Repairs": function.get("repairs", 0),
                    "Language": function.get("lang"),
                    "Total Binary Success": success,
                    "Total Binary Failed": failed,
                    "Total Repairs": repairs,
                })
            rows.append({
                "Run ID": run_id,
                "Binary": binary,
                "Function": "OVERALL",
                "Status": "Summary",
                "Repairs": repairs,
                "Language": "",
                "Total Binary Success": success,
                "Total Binary Failed": failed,
                "Total Repairs": repairs,
            })

        os.makedirs(os.path.dirname(os.path.abspath(self.file_path)), exist_ok=True)
        if os.path.exists(self.file_path):
            workbook = load_workbook(self.file_path)
        else:
            workbook = Workbook()
            if workbook.active and workbook.active.title == "Sheet":
                workbook.remove(workbook.active)

        if self.sheet_name in workbook.sheetnames:
            sheet = workbook[self.sheet_name]
        else:
            sheet = workbook.create_sheet(self.sheet_name)
            sheet.append(HEADERS)
        for row in rows:
            sheet.append([row.get(column) for column in HEADERS])
        sheet.freeze_panes = "A2"
        workbook.save(self.file_path)